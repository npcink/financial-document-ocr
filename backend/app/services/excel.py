import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional, Tuple
import os
import re
import csv
import io
import json
import zipfile
from datetime import datetime
from app.db import EXPORT_DIR
from app.utils.document_schema import document_type_label, field_type, ordered_columns_for, schema_for
from app.utils.field_normalizer import (
    EXPORT_CACHE_VERSION,
    extract_field_value,
    field_confidence,
    normalize_field_names,
    normalize_fields,
)
from app.utils.financial_fields import amount_value

BASE_HEADERS = ["文件名", "单据序号", "状态", "错误信息"]
DERIVED_HEADERS = ["唯一去重键", "数据完整性", "行复核状态", "金额状态", "时间状态", "收支方向", "绝对金额", "时间完整性"]
STANDARD_DERIVED_HEADERS = ["行复核状态", "金额状态", "时间状态", "收支方向", "绝对金额"]
FULL_DETAIL_HEADERS = ["文件名", "单据类型", "单据序号", "状态", "错误信息"] + DERIVED_HEADERS
AUDIT_HEADERS = ["其他字段JSON", "字段异常说明"]
STANDARD_AUDIT_HEADERS = ["字段异常说明"]
EXPORT_PROFILE_STANDARD = "standard"
EXPORT_PROFILE_FULL = "full"
STANDARD_EXPORT_COLUMNS = {
    "bank_receipt": [
        "付款方户名", "付款方账号", "收款方户名", "收款方账号",
        "交易金额", "交易日期", "时间戳", "交易流水号",
    ],
    "invoice": [
        "发票号码", "开票日期", "购买方名称", "销售方名称",
        "金额合计", "税额", "价税合计",
    ],
    "receipt": [
        "日期", "付款单位", "收款单位", "金额小写", "项目", "摘要",
    ],
    "transaction_list": [
        "交易类型", "交易对象", "交易金额", "交易日期", "时间戳",
    ],
    "subscription_record": [
        "产品名称", "订购状态", "订购时间", "有效期开始", "有效期结束", "退订时间", "月费",
    ],
    "unknown": [
        "日期", "交易日期", "时间戳", "交易时间",
        "金额", "交易金额", "金额小写", "价税合计",
        "交易对象", "付款方式", "交易状态",
        "产品名称", "商品名称", "摘要", "备注",
        "订单号", "商家订单号", "交易流水号",
    ],
    "general": [
        "日期", "交易日期", "时间戳", "交易时间",
        "金额", "交易金额", "金额小写", "价税合计",
        "交易对象", "付款方式", "交易状态",
        "产品名称", "商品名称", "摘要", "备注",
        "订单号", "商家订单号", "交易流水号",
    ],
}
LEGACY_DERIVED_HEADERS = ["识别质量"]
RESERVED_HEADERS = set(BASE_HEADERS + DERIVED_HEADERS + LEGACY_DERIVED_HEADERS + FULL_DETAIL_HEADERS + AUDIT_HEADERS + STANDARD_DERIVED_HEADERS + STANDARD_AUDIT_HEADERS)


def _sanitize_sheet_name(name: str) -> str:
    invalid = re.sub(r'[\[\]:*?/\\]', '', name)
    return invalid[:31] or 'Sheet'


def normalize_fields_for_export(fields: Any, doc_type: str = "") -> dict:
    normalized, _ = normalize_fields(fields, doc_type)
    return normalized


def normalize_field_names_for_export(field_names: Any, doc_type: str = "") -> list[str]:
    return normalize_field_names(field_names, doc_type)


def _coerce_excel_value(field_name: str, value: Any, doc_type: str = "") -> Tuple[Any, Optional[str]]:
    if value is None:
        return "", None
    text = str(value).strip()
    if not text:
        return "", None

    expected_type = field_type(field_name, doc_type)
    if expected_type == "amount" or any(token in field_name for token in ("金额", "合计", "税额", "小写")):
        parsed_amount = amount_value(text)
        if parsed_amount is not None:
            return float(parsed_amount), '#,##0.00'

    if expected_type == "datetime":
        normalized = text.replace("年", "-").replace("月", "-").replace("日", "")
        normalized = re.sub(r"\s+", " ", normalized).strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"):
            try:
                return datetime.strptime(normalized[:len("2026-05-23 12:34:56")], fmt), "yyyy-mm-dd hh:mm:ss"
            except ValueError:
                pass
        date_only = _date_from_value(text)
        if date_only:
            return date_only, None

    if expected_type == "date" or any(token in field_name for token in ("日期", "开票日期")):
        date_match = re.search(r"(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})", text)
        if date_match:
            year, month, day = date_match.groups()
            try:
                return datetime(int(year), int(month), int(day)), "yyyy-mm-dd"
            except ValueError:
                pass
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
            sample = text[:10] if "-" in fmt or "/" in fmt else text
            try:
                return datetime.strptime(sample, fmt), "yyyy-mm-dd"
            except ValueError:
                pass

    return value, None


def _date_from_value(value: Any) -> Any:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})", text)
    if not match:
        return None
    year, month, day = match.groups()
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def _join_date_and_time(date_value: Any, time_value: Any) -> Any:
    if date_value is None or time_value is None:
        return time_value
    date_text = str(date_value).strip()
    time_text = str(time_value).strip()
    if not date_text or not time_text:
        return time_value
    if re.search(r"\d{4}[年/-]\d{1,2}[月/-]\d{1,2}", time_text):
        return time_value
    time_match = re.search(r"(\d{1,2})[:：](\d{1,2})(?:[:：](\d{1,2}))?", time_text)
    if not time_match:
        return time_value
    date_match = re.search(r"(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})", date_text)
    if not date_match:
        return time_value
    year, month, day = date_match.groups()
    hour, minute, second = time_match.groups()
    second = second or "00"
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d} {int(hour):02d}:{int(minute):02d}:{int(second):02d}"


def _type_label(expected_type: str) -> str:
    return {
        "amount": "金额",
        "date": "日期",
        "datetime": "日期时间",
        "bank_account": "银行账号",
        "identifier": "编号",
    }.get(expected_type, expected_type or "文本")


def _field_value_matches_type(field_name: str, value: Any, doc_type: str = "") -> bool:
    if value is None:
        return True
    text = str(value).strip()
    if not text:
        return True

    expected_type = field_type(field_name, doc_type)
    if expected_type == "text":
        return True
    if expected_type == "amount":
        if "%" in text or "％" in text:
            return False
        return amount_value(text) is not None
    if expected_type == "date":
        return bool(re.search(r"\d{4}[年/-]\d{1,2}[月/-]\d{1,2}", text))
    if expected_type == "datetime":
        return bool(
            re.search(r"\d{4}[年/-]\d{1,2}[月/-]\d{1,2}", text)
            or re.search(r"\d{1,2}[:：]\d{1,2}", text)
        )
    if expected_type == "bank_account":
        digits = re.sub(r"\D", "", text)
        return 8 <= len(digits) <= 30
    if expected_type == "identifier":
        compact = re.sub(r"\s", "", text)
        return len(compact) >= 4
    return True


def _json_dumps(value: Any) -> str:
    if not value:
        return ""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _clean_amount(value: Any) -> Optional[float]:
    parsed = amount_value(value)
    if parsed is None:
        return None
    return float(parsed)


def _string_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _first_row_value(row: dict, names: list[str]) -> Any:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip() != "":
            return value
    return ""


def _first_field_value(fields: dict, names: list[str]) -> Any:
    for name in names:
        if name not in fields:
            continue
        value = extract_field_value(fields.get(name))
        if value is not None and str(value).strip() != "":
            return value
    return ""


def _dedupe_key(row: dict, doc_type: str) -> str:
    if doc_type == "bank_receipt":
        names = ["交易日期", "交易金额", "交易流水号", "付款方户名", "收款方户名"]
    elif doc_type == "invoice":
        names = ["发票代码", "发票号码", "开票日期", "价税合计", "购买方名称", "销售方名称"]
    elif doc_type == "receipt":
        names = ["收据编号", "日期", "金额小写", "付款单位", "收款单位"]
    elif doc_type == "transaction_list":
        names = ["交易日期", "时间戳", "交易金额", "交易类型", "交易对象", "原始时间文本"]
    elif doc_type == "subscription_record":
        names = ["产品名称", "订购时间", "有效期开始", "有效期结束", "退订时间", "月费"]
    else:
        names = ["交易日期", "日期", "交易金额", "金额", "交易流水号", "流水号"]
    parts = [_string_value(row.get(name)) for name in names]
    return "|".join(part for part in parts if part)


def _data_integrity(row: dict, doc_type: str) -> str:
    schema = schema_for(doc_type)
    required = schema.get("required", [])
    recommended = schema.get("recommended", [])
    missing_required = [name for name in required if not _string_value(row.get(name))]
    missing_recommended = [name for name in recommended if not _string_value(row.get(name))]
    if not missing_required and not missing_recommended:
        return "完整"
    parts = [f"缺{name}" for name in missing_required]
    parts.extend([f"缺{name}" for name in missing_recommended])
    return "；".join(parts)


def _amount_helpers(row: dict, doc_type: str) -> tuple[str, Any]:
    if doc_type == "subscription_record":
        return "", ""
    amount = _clean_amount(_first_row_value(row, _business_amount_fields(doc_type)))
    if amount is None:
        return "", ""
    if amount < 0:
        direction = "支出"
    elif amount > 0:
        direction = "收入"
    else:
        direction = "零金额"
    return direction, abs(amount)


def _business_amount_fields(doc_type: str) -> list[str]:
    if doc_type == "invoice":
        return ["价税合计", "价税合计小写", "金额合计", "交易金额", "金额"]
    if doc_type == "receipt":
        return ["金额小写", "金额", "交易金额"]
    if doc_type == "subscription_record":
        return ["月费", "金额", "费用"]
    return ["交易金额", "金额", "付款金额", "转账金额", "金额小写", "价税合计"]


def _active_amount_field(row: dict, doc_type: str) -> str:
    for name in _business_amount_fields(doc_type):
        if _clean_amount(row.get(name)) is not None:
            return name
    return ""


def _amount_status(row: dict, doc_type: str, low_conf_fields: Optional[List[str]] = None) -> str:
    field_name = _active_amount_field(row, doc_type)
    amount_required_types = {"bank_receipt", "transaction_list", "receipt", "invoice", "unknown", "general"}
    if not field_name:
        return "缺金额" if doc_type in amount_required_types else "无金额"
    if field_name in set(low_conf_fields or []):
        return "金额需复核"
    return "金额可靠"


def _has_date(value: Any) -> bool:
    return bool(re.search(r"\d{4}[年/-]\d{1,2}[月/-]\d{1,2}", _string_value(value)))


def _has_time(value: Any) -> bool:
    return bool(re.search(r"\d{1,2}[:：]\d{1,2}", _string_value(value)))


def _has_partial_month_day(value: Any) -> bool:
    return bool(re.search(r"\b\d{1,2}[-/]\d{1,2}(?:\s+\d{1,2}[:：]\d{1,2})?", _string_value(value)))


def _time_completeness(date_value: Any, original_timestamp: Any, final_timestamp: Any) -> str:
    if _has_date(original_timestamp) and _has_time(original_timestamp):
        return "完整时间"
    if _has_time(original_timestamp) and not _has_date(original_timestamp) and _has_date(date_value):
        return "仅时间补全"
    if _has_date(original_timestamp) and not _has_time(original_timestamp):
        return "仅日期"
    if _has_date(final_timestamp) and _has_time(final_timestamp):
        return "完整时间"
    if _has_date(date_value):
        return "仅日期"
    if _has_partial_month_day(date_value) or _has_partial_month_day(original_timestamp) or _has_partial_month_day(final_timestamp):
        return "缺年份"
    if _has_time(original_timestamp):
        return "仅时间"
    return "缺时间"


def _time_status(row: dict) -> str:
    time_integrity = _string_value(row.get("时间完整性"))
    return {
        "完整时间": "时间完整",
        "仅时间补全": "已补全年月日",
        "仅日期": "只有日期",
        "缺年份": "缺年份",
        "仅时间": "缺日期",
        "缺时间": "缺时间",
    }.get(time_integrity, time_integrity or "缺时间")


def _apply_derived_fields(
    row: dict,
    doc_type: str,
    low_conf_fields: Optional[List[str]] = None,
    original_date: Any = None,
    original_timestamp: Any = None,
) -> None:
    direction, absolute_amount = _amount_helpers(row, doc_type)
    row["收支方向"] = direction
    row["绝对金额"] = absolute_amount
    row["数据完整性"] = _data_integrity(row, doc_type)
    row["唯一去重键"] = _dedupe_key(row, doc_type)
    row["时间完整性"] = _time_completeness(
        row.get("交易日期") or row.get("日期") or original_date,
        original_timestamp if original_timestamp is not None else row.get("时间戳"),
        row.get("时间戳"),
    )
    row["金额状态"] = _amount_status(row, doc_type, low_conf_fields)
    row["时间状态"] = _time_status(row)
    row["行复核状态"] = _row_review_status(row, doc_type)
    row["识别质量"] = row["行复核状态"]


def _row_review_status(row: dict, doc_type: str) -> str:
    if row.get("状态") and row.get("状态") != "成功":
        return "处理失败"

    data_integrity = _string_value(row.get("数据完整性"))
    amount_status = _string_value(row.get("金额状态"))
    time_status = _string_value(row.get("时间状态"))

    if amount_status == "缺金额":
        return "缺金额"
    if "缺交易金额" in data_integrity or "缺金额小写" in data_integrity or "缺价税合计" in data_integrity:
        return "缺金额"
    if amount_status == "金额需复核":
        return "金额需复核"
    if doc_type in {"bank_receipt", "transaction_list", "unknown", "general"} and time_status in {"缺年份", "缺日期", "缺时间"}:
        return time_status
    if data_integrity == "完整":
        if doc_type in {"invoice", "receipt", "subscription_record"}:
            return "可直接使用"
        if time_status in {"时间完整", "已补全年月日", "只有日期"}:
            return "可直接使用"
    if data_integrity:
        return data_integrity
    return "需补充信息"


def _has_business_amount(row: dict, doc_type: str) -> bool:
    if doc_type == "invoice":
        names = ["价税合计", "价税合计小写", "金额合计"]
    elif doc_type == "receipt":
        names = ["金额小写", "金额"]
    elif doc_type == "subscription_record":
        names = ["月费"]
    else:
        names = ["交易金额", "金额", "付款金额", "转账金额"]
    return _clean_amount(_first_row_value(row, names)) is not None


def _compact_label(value: Any) -> str:
    return re.sub(r"[\s:：/／_-]+", "", str(value or "").strip())


def _looks_like_header_echo(field_name: str, value: Any, doc_type: str, columns: List[str]) -> bool:
    text = str(value or "").strip()
    if not text:
        return False

    schema = schema_for(doc_type)
    labels = set(BASE_HEADERS + DERIVED_HEADERS + AUDIT_HEADERS + columns + list(schema.get("order", [])))
    labels.update(schema.get("aliases", {}).keys())
    labels.update(schema.get("aliases", {}).values())
    labels.add(field_name)

    compact_text = _compact_label(text)
    return bool(compact_text) and compact_text in {_compact_label(label) for label in labels if label}


def _should_keep_low_confidence_marker(field_name: str, fields: dict, doc_type: str) -> bool:
    if field_name not in fields:
        return True
    value = extract_field_value(fields.get(field_name))
    if _clean_amount(value) is not None and field_name in _business_amount_fields(doc_type):
        return field_confidence(fields.get(field_name)) < 0.8
    return True


def _export_low_confidence_fields(low_conf_fields: list[str], fields: dict, doc_type: str) -> list[str]:
    return [
        field_name for field_name in low_conf_fields
        if _should_keep_low_confidence_marker(field_name, fields, doc_type)
    ]


def _normalize_export_profile(export_profile: str = EXPORT_PROFILE_STANDARD) -> str:
    return EXPORT_PROFILE_FULL if export_profile == EXPORT_PROFILE_FULL else EXPORT_PROFILE_STANDARD


def derived_headers_for(export_profile: str = EXPORT_PROFILE_STANDARD) -> list[str]:
    if _normalize_export_profile(export_profile) == EXPORT_PROFILE_FULL:
        return DERIVED_HEADERS
    return STANDARD_DERIVED_HEADERS


def audit_headers_for(export_profile: str = EXPORT_PROFILE_STANDARD) -> list[str]:
    if _normalize_export_profile(export_profile) == EXPORT_PROFILE_FULL:
        return AUDIT_HEADERS
    return STANDARD_AUDIT_HEADERS


def export_columns_for(
    doc_type: str,
    all_fields: set[str],
    custom_columns: List[str] = None,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> list[str]:
    if not custom_columns and _normalize_export_profile(export_profile) == EXPORT_PROFILE_STANDARD:
        standard_columns = STANDARD_EXPORT_COLUMNS.get(doc_type)
        if standard_columns:
            return [column for column in standard_columns if column not in RESERVED_HEADERS]

    return [
        column for column in ordered_columns_for(doc_type, all_fields, custom_columns)
        if column not in RESERVED_HEADERS and column not in DERIVED_HEADERS
    ]


def _prepare_export_row(result: Dict[str, Any], columns: List[str]) -> tuple[dict, list[str], list[str]]:
    doc_type = result.get("document_type", "")
    fields, conflict_fields = normalize_fields(result.get("structured_data", {}), doc_type)
    low_conf_fields = normalize_field_names_for_export(
        result.get("low_confidence_fields", []),
        doc_type,
    )
    for field_name in conflict_fields:
        if field_name not in low_conf_fields:
            low_conf_fields.append(field_name)

    row = _base_result_row(result)
    known_columns = set(columns)
    extra_fields = {}
    issues = []
    invalid_fields = []

    for field_name, field_data in fields.items():
        value = extract_field_value(field_data)
        if field_name in RESERVED_HEADERS:
            if value is not None and str(value).strip():
                issues.append(f"{field_name}: 识别结果包含系统列，已忽略")
            continue

        if field_name not in known_columns:
            if value is not None and str(value).strip():
                extra_fields[field_name] = value
            continue

        if _looks_like_header_echo(field_name, value, doc_type, columns):
            issues.append(f"{field_name}: 识别结果为表头“{value}”，已置空")
            invalid_fields.append(field_name)
            continue

        if not _field_value_matches_type(field_name, value, doc_type):
            expected = _type_label(field_type(field_name, doc_type))
            issues.append(f"{field_name}: {value} 不符合{expected}格式，已移入异常说明")
            invalid_fields.append(field_name)
            continue

        row[field_name] = value

    original_date = row.get("交易日期") or row.get("日期") or _first_field_value(fields, ["交易日期", "日期"])
    original_timestamp = row.get("时间戳")

    if "交易日期" in known_columns and not row.get("交易日期") and row.get("时间戳"):
        inferred_date = _date_from_value(row.get("时间戳"))
        if inferred_date:
            row["交易日期"] = inferred_date
    if "时间戳" in known_columns and row.get("交易日期") and row.get("时间戳"):
        row["时间戳"] = _join_date_and_time(row.get("交易日期"), row.get("时间戳"))
    low_conf_fields = _export_low_confidence_fields(low_conf_fields, fields, doc_type)
    _apply_derived_fields(
        row,
        doc_type,
        low_conf_fields=low_conf_fields,
        original_date=original_date,
        original_timestamp=original_timestamp,
    )

    if extra_fields:
        row["其他字段JSON"] = _json_dumps(extra_fields)
    if issues:
        row["字段异常说明"] = "；".join(issues)

    low_conf_fields.extend([field for field in invalid_fields if field not in low_conf_fields])
    return row, low_conf_fields, invalid_fields


def _status_label(status: str) -> str:
    return {
        "completed": "成功",
        "failed": "失败",
        "pending": "等待处理",
    }.get(status or "", status or "")


def _base_result_row(result: Dict[str, Any]) -> dict:
    document_index = result.get("source_document_index") or 1
    document_count = result.get("source_document_count") or 1
    document_label = f"{document_index}/{document_count}" if document_count > 1 else "1"
    return {
        "文件名": result.get("filename", ""),
        "单据序号": document_label,
        "状态": _status_label(result.get("status")),
        "错误信息": result.get("error_message") or "",
    }


def flatten_result_row(result: Dict[str, Any], columns: List[str]) -> dict:
    row, _, _ = _prepare_export_row(result, columns)
    return row


def export_headers_for(
    doc_type: str,
    all_fields: set[str],
    custom_columns: List[str] = None,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> list[str]:
    columns = export_columns_for(doc_type, all_fields, custom_columns, export_profile)
    return BASE_HEADERS + derived_headers_for(export_profile) + columns + audit_headers_for(export_profile)


def _sheet_title_for(doc_type: str, used_titles: set[str]) -> str:
    base = _sanitize_sheet_name(document_type_label(doc_type))
    title = base
    index = 2
    while title in used_titles:
        suffix = f"_{index}"
        title = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used_titles.add(title)
    return title


def _full_field_column_name(field_name: str) -> str:
    return f"识别字段:{field_name}" if field_name in set(FULL_DETAIL_HEADERS + AUDIT_HEADERS) else field_name


def _collect_full_field_columns(results: List[Dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen = set()
    extras_by_type: dict[str, set[str]] = {}
    doc_types = []

    for result in results:
        doc_type = result.get("document_type", "unknown")
        if doc_type not in doc_types:
            doc_types.append(doc_type)
        fields = normalize_fields_for_export(result.get("structured_data", {}), doc_type)
        extras_by_type.setdefault(doc_type, set()).update(fields.keys())

    for doc_type in doc_types:
        schema_order = schema_for(doc_type).get("order", [])
        fields = extras_by_type.get(doc_type, set())
        ordered = [field for field in schema_order if field in fields]
        ordered.extend(sorted(fields - set(ordered)))
        for field_name in ordered:
            column_name = _full_field_column_name(field_name)
            if column_name not in seen:
                columns.append(column_name)
                seen.add(column_name)

    return columns


def _full_detail_row(result: Dict[str, Any], field_columns: List[str]) -> dict:
    doc_type = result.get("document_type", "unknown")
    base = _base_result_row(result)
    row = {
        "文件名": base["文件名"],
        "单据类型": document_type_label(doc_type),
        "单据序号": base["单据序号"],
        "状态": base["状态"],
        "错误信息": base["错误信息"],
    }
    fields = normalize_fields_for_export(result.get("structured_data", {}), doc_type)
    for field_name, field_data in fields.items():
        column_name = _full_field_column_name(field_name)
        if column_name in field_columns:
            row[column_name] = extract_field_value(field_data)
    original_date = row.get("交易日期")
    original_timestamp = row.get("时间戳")
    if "交易日期" in field_columns and not row.get("交易日期") and row.get("时间戳"):
        inferred_date = _date_from_value(row.get("时间戳"))
        if inferred_date:
            row["交易日期"] = inferred_date
    if "时间戳" in field_columns and row.get("交易日期") and row.get("时间戳"):
        row["时间戳"] = _join_date_and_time(row.get("交易日期"), row.get("时间戳"))
    _apply_derived_fields(
        row,
        doc_type,
        low_conf_fields=[],
        original_date=original_date,
        original_timestamp=original_timestamp,
    )
    return row


def _create_full_detail_sheet(
    wb,
    results: List[Dict[str, Any]],
    used_titles: set[str],
    header_font,
    header_fill,
    header_alignment,
    cell_font,
    cell_alignment,
    thin_border,
) -> None:
    if not results:
        return
    ws = wb.create_sheet(title=_sheet_title_for("全字段明细", used_titles))
    field_columns = _collect_full_field_columns(results)
    headers = FULL_DETAIL_HEADERS + field_columns

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, result in enumerate(results, 2):
        row = _full_detail_row(result, field_columns)
        doc_type = result.get("document_type", "unknown")
        for col_idx, header in enumerate(headers, 1):
            value = row.get(header, "")
            if header not in FULL_DETAIL_HEADERS:
                source_field = header.replace("识别字段:", "", 1)
                value, number_format = _coerce_excel_value(source_field, value, doc_type)
            else:
                number_format = None
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border
            if number_format:
                cell.number_format = number_format

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, len(headers) + 1):
        header = str(ws.cell(row=1, column=col_idx).value or "")
        max_len = len(header)
        for row_idx in range(2, min(ws.max_row, 200) + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(42, max_len + 2))
    ws.column_dimensions["A"].width = 30


def _append_full_detail_write_only_sheet(
    wb,
    results: List[Dict[str, Any]],
    used_titles: set[str],
    header_font,
    header_fill,
    header_alignment,
    cell_font,
    cell_alignment,
    thin_border,
) -> None:
    if not results:
        return
    ws = wb.create_sheet(title=_sheet_title_for("全字段明细", used_titles))
    field_columns = _collect_full_field_columns(results)
    headers = FULL_DETAIL_HEADERS + field_columns
    ws.append([
        _styled_write_only_cell(
            ws,
            header,
            font=header_font,
            fill=header_fill,
            alignment=header_alignment,
            border=thin_border,
        )
        for header in headers
    ])

    for result in results:
        row = _full_detail_row(result, field_columns)
        doc_type = result.get("document_type", "unknown")
        cells = []
        for header in headers:
            value = row.get(header, "")
            if header not in FULL_DETAIL_HEADERS:
                source_field = header.replace("识别字段:", "", 1)
                value, number_format = _coerce_excel_value(source_field, value, doc_type)
            else:
                number_format = None
            cells.append(_styled_write_only_cell(
                ws,
                value,
                font=cell_font,
                alignment=cell_alignment,
                border=thin_border,
                number_format=number_format,
            ))
        ws.append(cells)


def csv_line(values: list[Any]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["" if value is None else value for value in values])
    return buffer.getvalue()


def csv_content_for_results(
    doc_type: str,
    results: List[Dict[str, Any]],
    custom_columns: List[str] = None,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> str:
    all_fields = set()
    for result in results:
        fields = normalize_fields_for_export(result.get("structured_data", {}), doc_type)
        if isinstance(fields, dict):
            all_fields.update(fields.keys())
    columns = export_columns_for(doc_type, all_fields, custom_columns, export_profile)
    headers = BASE_HEADERS + derived_headers_for(export_profile) + columns + audit_headers_for(export_profile)

    chunks = ["\ufeff", csv_line(headers)]
    for result in results:
        row = flatten_result_row(result, columns)
        chunks.append(csv_line([row.get(header, "") for header in headers]))
    return "".join(chunks)


def generate_csv_zip(
    job_id: str,
    results: List[Dict[str, Any]],
    custom_columns: List[str] = None,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> str:
    export_dir = os.path.join(EXPORT_DIR, job_id)
    os.makedirs(export_dir, exist_ok=True)
    export_profile = _normalize_export_profile(export_profile)
    export_path = os.path.join(export_dir, f"result_{job_id}_{export_profile}.zip")

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for result in results:
        doc_type = result.get("document_type", "unknown")
        grouped.setdefault(doc_type, []).append(result)

    used_names = set()
    with zipfile.ZipFile(export_path, "w", compression=zipfile.ZIP_DEFLATED) as zip_file:
        for doc_type, items in grouped.items():
            base = _sanitize_sheet_name(document_type_label(doc_type))
            filename = f"{base}.csv"
            index = 2
            while filename in used_names:
                filename = f"{base}_{index}.csv"
                index += 1
            used_names.add(filename)
            zip_file.writestr(filename, csv_content_for_results(doc_type, items, custom_columns, export_profile))

    return export_path


def generate_excel(
    job_id: str,
    results: List[Dict[str, Any]],
    custom_columns: List[str] = None,
    updated_at: str = None,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> str:
    export_dir = os.path.join(EXPORT_DIR, job_id)
    os.makedirs(export_dir, exist_ok=True)
    export_profile = _normalize_export_profile(export_profile)
    export_path = os.path.join(export_dir, f"result_{job_id}_{export_profile}.xlsx")
    meta_path = os.path.join(export_dir, f"result_{job_id}_{export_profile}.export.json")

    if os.path.exists(export_path) and updated_at:
        cache_mtime = os.path.getmtime(export_path)
        try:
            from datetime import datetime
            updated_dt = datetime.fromisoformat(updated_at)
            meta = _read_export_meta(meta_path)
            if (
                cache_mtime > updated_dt.timestamp()
                and cache_mtime > _export_dependency_mtime()
                and meta.get("version") == EXPORT_CACHE_VERSION
            ):
                return export_path
        except (ValueError, OSError):
            pass

    if len(results) > 5000:
        return _generate_excel_write_only(job_id, results, custom_columns, export_path, meta_path, export_profile)

    wb = openpyxl.Workbook()

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for result in results:
        doc_type = result.get("document_type", "unknown")
        grouped.setdefault(doc_type, []).append(result)

    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    cell_font = Font(name="微软雅黑", size=10)
    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    low_conf_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    low_conf_font = Font(name="微软雅黑", size=10, color="9C0006")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    first_sheet = True
    used_titles = set()
    for doc_type, items in grouped.items():
        if first_sheet:
            ws = wb.active
            first_sheet = False
        else:
            ws = wb.create_sheet()

        ws.title = _sheet_title_for(doc_type, used_titles)

        all_fields = set()
        for result in items:
            fields = normalize_fields_for_export(result.get("structured_data", {}), doc_type)
            if isinstance(fields, dict):
                all_fields.update(fields.keys())

        columns = export_columns_for(doc_type, all_fields, custom_columns, export_profile)

        derived_headers = derived_headers_for(export_profile)
        audit_headers = audit_headers_for(export_profile)
        headers = BASE_HEADERS + derived_headers + columns + audit_headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for row_idx, result in enumerate(items, 2):
            ws.cell(row=row_idx, column=1, value=result.get("filename", "")).font = cell_font
            ws.cell(row=row_idx, column=1).border = thin_border
            ws.cell(row=row_idx, column=1).alignment = cell_alignment
            document_index = result.get("source_document_index") or 1
            document_count = result.get("source_document_count") or 1
            document_label = f"{document_index}/{document_count}" if document_count > 1 else "1"
            base_values = [
                result.get("filename", ""),
                document_label,
                _status_label(result.get("status")),
                result.get("error_message") or "",
            ]
            for base_col_idx, value in enumerate(base_values, 1):
                cell = ws.cell(row=row_idx, column=base_col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border

            export_row, low_conf_fields, invalid_fields = _prepare_export_row(result, columns)

            for derived_col_idx, derived_name in enumerate(derived_headers, len(BASE_HEADERS) + 1):
                value = export_row.get(derived_name, "")
                value, number_format = _coerce_excel_value(derived_name, value, doc_type)
                cell = ws.cell(row=row_idx, column=derived_col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border
                if number_format:
                    cell.number_format = number_format

            for col_idx, col_name in enumerate(columns, len(BASE_HEADERS) + len(derived_headers) + 1):
                value = export_row.get(col_name, "")
                value, number_format = _coerce_excel_value(col_name, value, doc_type)

                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border
                if number_format:
                    cell.number_format = number_format

                if col_name in low_conf_fields:
                    cell.fill = low_conf_fill
                    cell.font = low_conf_font

            audit_start = len(BASE_HEADERS) + len(derived_headers) + len(columns) + 1
            for audit_idx, audit_name in enumerate(audit_headers, audit_start):
                cell = ws.cell(row=row_idx, column=audit_idx, value=export_row.get(audit_name, ""))
                cell.font = low_conf_font if audit_name == "字段异常说明" and invalid_fields else cell_font
                if audit_name == "字段异常说明" and invalid_fields:
                    cell.fill = low_conf_fill
                cell.alignment = cell_alignment
                cell.border = thin_border

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, len(headers) + 1):
            header = str(ws.cell(row=1, column=col_idx).value or "")
            max_len = len(header)
            for row_idx in range(2, min(ws.max_row, 200) + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                max_len = max(max_len, len(str(value)) if value is not None else 0)
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(42, max_len + 2))

        ws.column_dimensions["A"].width = 30

    if export_profile == EXPORT_PROFILE_FULL:
        _create_full_detail_sheet(
            wb,
            results,
            used_titles,
            header_font,
            header_fill,
            header_alignment,
            cell_font,
            cell_alignment,
            thin_border,
        )

    wb.save(export_path)
    _write_export_meta(meta_path)

    return export_path


def _styled_write_only_cell(ws, value: Any, font=None, fill=None, alignment=None, border=None, number_format: str = None):
    cell = WriteOnlyCell(ws, value=value)
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


def _generate_excel_write_only(
    job_id: str,
    results: List[Dict[str, Any]],
    custom_columns: List[str],
    export_path: str,
    meta_path: str,
    export_profile: str = EXPORT_PROFILE_STANDARD,
) -> str:
    wb = openpyxl.Workbook(write_only=True)

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for result in results:
        doc_type = result.get("document_type", "unknown")
        grouped.setdefault(doc_type, []).append(result)

    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    cell_font = Font(name="微软雅黑", size=10)
    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    low_conf_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    low_conf_font = Font(name="微软雅黑", size=10, color="9C0006")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    used_titles = set()
    for doc_type, items in grouped.items():
        ws = wb.create_sheet(title=_sheet_title_for(doc_type, used_titles))
        all_fields = set()
        for result in items:
            fields = normalize_fields_for_export(result.get("structured_data", {}), doc_type)
            if isinstance(fields, dict):
                all_fields.update(fields.keys())
        columns = export_columns_for(doc_type, all_fields, custom_columns, export_profile)
        derived_headers = derived_headers_for(export_profile)
        audit_headers = audit_headers_for(export_profile)
        headers = BASE_HEADERS + derived_headers + columns + audit_headers
        ws.append([
            _styled_write_only_cell(
                ws,
                header,
                font=header_font,
                fill=header_fill,
                alignment=header_alignment,
                border=thin_border,
            )
            for header in headers
        ])

        for result in items:
            row_cells = []
            export_row, low_conf_fields, invalid_fields = _prepare_export_row(result, columns)

            for header in BASE_HEADERS + derived_headers:
                row_cells.append(_styled_write_only_cell(
                    ws,
                    export_row.get(header, ""),
                    font=cell_font,
                    alignment=cell_alignment,
                    border=thin_border,
                ))
            for col_name in columns:
                value = export_row.get(col_name, "")
                value, number_format = _coerce_excel_value(col_name, value, doc_type)
                row_cells.append(_styled_write_only_cell(
                    ws,
                    value,
                    font=low_conf_font if col_name in low_conf_fields else cell_font,
                    fill=low_conf_fill if col_name in low_conf_fields else None,
                    alignment=cell_alignment,
                    border=thin_border,
                    number_format=number_format,
                ))
            for audit_name in audit_headers:
                row_cells.append(_styled_write_only_cell(
                    ws,
                    export_row.get(audit_name, ""),
                    font=low_conf_font if audit_name == "字段异常说明" and invalid_fields else cell_font,
                    fill=low_conf_fill if audit_name == "字段异常说明" and invalid_fields else None,
                    alignment=cell_alignment,
                    border=thin_border,
                ))
            ws.append(row_cells)

    if _normalize_export_profile(export_profile) == EXPORT_PROFILE_FULL:
        _append_full_detail_write_only_sheet(
            wb,
            results,
            used_titles,
            header_font,
            header_fill,
            header_alignment,
            cell_font,
            cell_alignment,
            thin_border,
        )

    wb.save(export_path)
    _write_export_meta(meta_path)
    return export_path


def _export_dependency_mtime() -> float:
    import app.utils.document_schema as document_schema
    import app.utils.field_normalizer as field_normalizer
    import app.utils.financial_fields as financial_fields
    import app.utils.ocr_preprocessor as ocr_preprocessor

    return max(
        os.path.getmtime(__file__),
        os.path.getmtime(document_schema.__file__),
        os.path.getmtime(field_normalizer.__file__),
        os.path.getmtime(financial_fields.__file__),
        os.path.getmtime(ocr_preprocessor.__file__),
    )


def _read_export_meta(meta_path: str) -> dict:
    try:
        import json
        with open(meta_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def _write_export_meta(meta_path: str) -> None:
    import json
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump({"version": EXPORT_CACHE_VERSION}, f, ensure_ascii=False)
